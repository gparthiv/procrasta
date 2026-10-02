from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from pptx import Presentation
from PIL import Image
from google import genai
from google.genai import errors
from fastapi import HTTPException


def parse_pdf(
    file_path: str,
    filename: str,
):
    loader = PyPDFLoader(file_path)

    raw_docs = loader.load()

    extracted_docs = []

    for doc in raw_docs:
        extracted_docs.append(
            Document(
                page_content=doc.page_content,
                metadata={
                    "source": filename,
                    "page": doc.metadata.get("page", 0) + 1,
                    "type": "pdf",
                },
            )
        )

    return extracted_docs


def parse_pptx(
    file_path: str,
    filename: str,
):
    prs = Presentation(file_path)

    docs = []

    for slide_idx, slide in enumerate(
        prs.slides,
        start=1,
    ):

        text_runs = []

        for shape in slide.shapes:

            if shape.has_text_frame:

                for paragraph in shape.text_frame.paragraphs:
                    text_runs.append(paragraph.text)

        slide_text = "\n".join(text_runs).strip()

        if slide_text:

            docs.append(
                Document(
                    page_content=slide_text,
                    metadata={
                        "source": filename,
                        "slide": slide_idx,
                        "type": "pptx",
                    },
                )
            )

    return docs


def ocr_image_with_gemini(
    image_path: str,
    filename: str,
    client,
):
    """Extract text from handwritten or scanned study material using Gemini."""

    img = Image.open(image_path)

    prompt = """
Transcribe all visible text from this study document accurately.

Requirements:
- Preserve headings and structure where possible.
- Preserve mathematical formulas and symbols.
- Preserve bullet points and lists.
- Do not summarize.
- Do not add information that is not visible.
- Return only the extracted/transcribed text.
"""

    try:

        response = client.models.generate_content(
            model="gemini-3.5-flash",
            contents=[prompt, img],
        )

    except errors.APIError as e:

        raise HTTPException(
            status_code=503,
            detail=f"Gemini OCR temporarily unavailable: {e}",
        )

    return [
        Document(
            page_content=response.text,
            metadata={
                "source": filename,
                "page": 1,
                "type": "handwritten",
            },
        )
    ]
