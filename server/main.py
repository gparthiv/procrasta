import os
import tempfile
from fastapi import FastAPI, UploadFile, File, HTTPException
from langchain_community.document_loaders import PyPDFLoader
from dotenv import load_dotenv
from langchain_google_genai import GoogleGenerativeAIEmbeddings, ChatGoogleGenerativeAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation
from google import genai
from PIL import Image
from google.genai import errors
from langchain_core.documents import Document
from qdrant_client import QdrantClient
from qdrant_client.http.models import Distance, VectorParams
from langchain_qdrant import QdrantVectorStore
from pydantic import BaseModel, Field

load_dotenv()
api_key = os.getenv("GOOGLE_API_KEY")
if not api_key:
    raise ValueError("Google APIKEY is not there")

client = genai.Client(api_key=os.getenv("GOOGLE_API_KEY"))

app = FastAPI(title="ProCrasto")

qdrant_client = QdrantClient(url="http://localhost:6333")
COLLECTION_NAME = "study_notes"

if not qdrant_client.collection_exists(COLLECTION_NAME):
    qdrant_client.create_collection(
        collection_name=COLLECTION_NAME,
        vectors_config=VectorParams(
            size=768,
            distance=Distance.COSINE,
        ),
    )

llm = ChatGoogleGenerativeAI(
    model="gemini-3.5-flash",
    google_api_key=api_key,
    temperature=0.2,
)

# Gemini Embeddings setup
embeddings = GoogleGenerativeAIEmbeddings(
    model="gemini-embedding-2",
    output_dimensionality=768,
    google_api_key=api_key,
)

vector_store = QdrantVectorStore(
    client=qdrant_client, collection_name=COLLECTION_NAME, embedding=embeddings
)


# HELPER FUNCTIONS LIKE OCR AND PPT PARSER
# function to parse ppt slides
def parse_pptx(file_path: str, filename: str):
    prs = Presentation(file_path)
    docs = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        text_runs = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text_runs.append(paragraph.text)
        slide_text = "\n".join(text_runs).strip()
        if slide_text:
            docs.append(
                Document(
                    page_content=slide_text,
                    metadata={"source": filename, "slide": slide_idx, "type": "pptx"},
                )
            )
    return docs


# function to image OCR using gemini model and prompt
def ocr_image_with_gemini(image_path: str, filename: str):
    """Extract text from handwritten or scanned study material using Gemini."""

    img = Image.open(image_path)

    prompt = """
    Transcribe all visible text from this study document accurately.

    Requirements:
    - Preserve headings and structure where possible.
    - Preserve mathematical formulas and symbols.
    - Preserve bullet points and lists.
    - Do not summarize.
    - Do not add information that is not visible.
    - Return only the extracted/transcribed text.
    """

    try:
        response = client.models.generate_content(
            model="gemini-3.5-flash", contents=[prompt, img]
        )
    except errors.APIError as e:
        raise HTTPException(
            status_code=503, detail=f"Gemini OCR temporarily unavailable: {e}"
        )

    return [
        Document(
            page_content=response.text,
            metadata={"source": filename, "page": 1, "type": "handwritten"},
        )
    ]


@app.get("/health")
def health_check():
    return {"status": "healthy"}


MAX_FILE_SIZE = 20 * 1024 * 1024


@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    """Receives notes/PDF/PPTX/Images, parses, chunks, and inserts into Qdrant."""
    ext = os.path.splitext(file.filename)[1].lower()
    # Read file once
    contents = await file.read()
    if len(contents) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File too large. Maximum size is 20 MB.",
        )
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        tmp.write(contents)
        tmp_path = tmp.name

    extracted_docs = []
    try:
        if ext == ".pdf":
            loader = PyPDFLoader(tmp_path)
            raw_docs = loader.load()
            for doc in raw_docs:
                extracted_docs.append(
                    Document(
                        page_content=doc.page_content,
                        metadata={
                            "source": file.filename,
                            "page": doc.metadata.get("page", 0) + 1,
                            "type": "pdf",
                        },
                    )
                )
        elif ext == ".pptx":
            extracted_docs = parse_pptx(tmp_path, file.filename)

        elif ext in [".png", ".jpg", ".jpeg"]:
            extracted_docs = ocr_image_with_gemini(
                tmp_path,
                file.filename,
            )

        else:
            raise HTTPException(
                status_code=400,
                detail="Unsupported file format.",
            )

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

    # Remove documents where extraction produced no actual text
    extracted_docs = [
        doc for doc in extracted_docs if doc.page_content and doc.page_content.strip()
    ]

    # Nothing useful was extracted
    if not extracted_docs:
        raise HTTPException(
            status_code=400,
            detail="Could not extract any text from the file.",
        )

    # Split extracted text into chunks
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200,
    )

    chunks = splitter.split_documents(extracted_docs)

    # Safety check: don't insert empty chunks into Qdrant
    chunks = [
        chunk for chunk in chunks if chunk.page_content and chunk.page_content.strip()
    ]

    if not chunks:
        raise HTTPException(
            status_code=400,
            detail="No usable text was found after chunking.",
        )

    # Store chunks in Qdrant
    vector_store.add_documents(chunks)

    return {
        "status": "success",
        "filename": file.filename,
        "documents_extracted": len(extracted_docs),
        "chunks_created": len(chunks),
    }


@app.get("/test-embedding")
def test_embedding():
    text = "CPU scheduling is the process of selecting a process for execution."
    vector = embeddings.embed_query(text)
    return {
        "status": "success",
        "dimensions": len(vector),
        "first_10_values": vector[:10],
    }


class ChatRequest(BaseModel):
    query: str = Field(..., min_length=2)


@app.post("/search")
async def search_endpoint(req: ChatRequest):
    results = vector_store.similarity_search(req.query, k=3)

    return {
        "query": req.query,
        "results": [
            {
                "text": doc.page_content,
                "metadata": doc.metadata,
            }
            for doc in results
        ],
    }


import time


import time

@app.post("/chat")
async def chat_endpoint(req: ChatRequest):

    total_start = time.perf_counter()

    # 1. Retrieval
    start = time.perf_counter()

    search_results = vector_store.similarity_search(
        req.query,
        k=3
    )

    retrieval_time = time.perf_counter() - start

    # 2. Context construction
    start = time.perf_counter()

    context_blocks = []

    for x in search_results:
        source = x.metadata.get("source", "Unknown File")

        if x.metadata.get("type") == "pptx":
            location = f"Slide {x.metadata.get('slide', '?')}"
        else:
            location = f"Page {x.metadata.get('page', '?')}"

        context_blocks.append(
            f"Source: {source}\n"
            f"Location: {location}\n"
            f"{x.page_content}"
        )

    context_str = "\n\n".join(context_blocks)

    context_time = time.perf_counter() - start

    # 3. Prompt construction
    start = time.perf_counter()

    prompt = f"""
You are an AI Study Companion.

Answer ONLY from the provided context.
Do not use outside knowledge or invent information.

If the answer is not present, say:
"I couldn't find this information in your uploaded notes."

Give a concise answer and cite the relevant source and page/slide.

Context:
{context_str}

Question:
{req.query}
"""

    prompt_time = time.perf_counter() - start

    # 4. Gemini
    start = time.perf_counter()

    response = llm.invoke(prompt)

    gemini_time = time.perf_counter() - start

    # 5. Total
    total_time = time.perf_counter() - total_start

    print("\n--- CHAT TIMING ---")
    print(f"Retrieval : {retrieval_time:.3f}s")
    print(f"Context   : {context_time:.3f}s")
    print(f"Prompt    : {prompt_time:.3f}s")
    print(f"Gemini    : {gemini_time:.3f}s")
    print(f"TOTAL     : {total_time:.3f}s")
    print("-------------------\n")

    return {
        "answer": response.content
    }